"""CLI entrypoint: JSON section list -> .docx report.

Invoked directly as a script (not via -m), so it adds its own directory to
sys.path to resolve the sibling `sections`/`charts`/`theme` modules
regardless of the caller's working directory.

Shares the section-JSON contract (and map_report.py) with render-ppt:
sections of type kpi_cards / table / chart / text / heading, plus
analysis_slot placeholders filled from --analysis.
"""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402

import sections as sec  # noqa: E402


def _merge_headings(sections):
    """A standalone heading followed by a heading-less section becomes that
    section's heading instead of floating as a bare paragraph."""
    merged = []
    pending = None
    for section in sections:
        if section["type"] == "heading":
            if pending is not None:
                merged.append(pending)
            pending = section
            continue
        if pending is not None:
            if not section.get("heading"):
                section = {**section, "heading": pending["text"]}
            else:
                merged.append(pending)
            pending = None
        merged.append(section)
    if pending is not None:
        merged.append(pending)
    return merged


def _fill_analysis_slots(sections, analysis):
    filled = []
    for section in sections:
        if section.get("type") != "analysis_slot":
            filled.append(section)
            continue
        entry = (analysis or {}).get(section["key"])
        if entry and entry.get("body"):
            filled.append({"type": "text",
                           "heading": entry.get("heading", section.get("heading")),
                           "body": entry["body"]})
        else:
            filled.append({"type": "text", "heading": section.get("heading"),
                           "body": "데이터 준비 중"})
    return filled


def build_document(data, analysis=None):
    document = Document()
    sec.setup_document(document)
    sec.add_title(document, data["title"], data.get("period"))
    sections = _fill_analysis_slots(data["sections"], analysis)
    heading_no = 0  # numbered section banners: 01, 02, ...
    for section in _merge_headings(sections):
        stype = section["type"]
        # headings render centrally so every banner gets a sequence number
        heading = section["text"] if stype == "heading" else section.get("heading")
        if heading:
            heading_no += 1
            # every section opens on a fresh page, except the first one
            # (which shares the title page)
            sec.add_heading(document, heading, f"{heading_no:02d}",
                            page_break=heading_no > 1)
        if stype == "heading":
            continue
        if stype == "kpi_cards":
            sec.add_kpi_cards(document, section["cards"])
        elif stype == "table":
            sec.add_data_table(document, None, section["headers"], section["rows"],
                               rows_total=section.get("rows_total"),
                               context=heading)
        elif stype == "chart":
            sec.add_combo_chart_section(
                document, None, section["categories"],
                section["bar_series"], section["line_series"])
        elif stype == "line_chart":
            sec.add_line_chart_section(
                document, None, section["categories"], section["series"])
        elif stype == "text":
            sec.add_text_section(document, None, section["body"])
        else:
            raise ValueError(f"unknown section type: {stype}")
    sec.add_footer(document)
    return document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", required=True, help="path to the section-list JSON file")
    parser.add_argument("--out", required=True, help="path to write the .docx file to")
    parser.add_argument("--analysis", default=None,
                        help="optional JSON filling map_report.py analysis_slot "
                             "placeholders: {\"section3\": {\"heading\": ..., \"body\": ...}}")
    args = parser.parse_args()

    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    analysis = None
    if args.analysis:
        analysis = json.loads(Path(args.analysis).read_text(encoding="utf-8"))
    document = build_document(data, analysis)

    out_path = Path(args.out).expanduser()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))


if __name__ == "__main__":
    main()
